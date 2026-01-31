import sqlite3
import google.generativeai as genai
import json
import os
import docx
import PyPDF2
import csv
import openpyxl


def guess_file_context(filename, content_preview):
    prompt = f"""Dưới đây là nội dung mẫu của file '{filename}':
{content_preview}
Hãy mô tả ngắn ý nghĩa hoặc chủ đề của file này bằng tiếng Việt."""
    model = genai.GenerativeModel("gemini-2.0-flash")
    response = model.generate_content(prompt)
    return response.text.strip()

def extract_preview_from_json(json_filename):
    with open(json_filename, "r", encoding="utf-8", errors="ignore") as f:
        data = json.load(f)
    return json.dumps(data, ensure_ascii=False, indent=2)[:500]

def extract_preview_from_xlsx(xlsx_filename):
    wb = openpyxl.load_workbook(xlsx_filename)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows.append(", ".join([str(cell) for cell in row]))
        if i >= 4: break
    return "\n".join(rows)

def extract_preview_from_docx(docx_filename):
    doc = docx.Document(docx_filename)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paras[:5])

def extract_preview_from_pdf(pdf_filename):
    text = ""
    with open(pdf_filename, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for i, page in enumerate(reader.pages):
            text += page.extract_text() or ""
            if i >= 1: break
    return text[:1000]

def extract_preview_from_csv(csv_filename):
    with open(csv_filename, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        rows = []
        for i, row in enumerate(reader):
            rows.append(", ".join(row))
            if i >= 4: break
    return "\n".join(rows)

def get_files_context(file_list, upload_dir, save_json=True):
    context_dict = {}
    for filename in file_list:
        ext = os.path.splitext(filename)[1].lower()
        file_path = os.path.join(upload_dir, filename)
        try:
            if ext == ".json":
                preview = extract_preview_from_json(file_path)
            elif ext == ".xlsx":
                preview = extract_preview_from_xlsx(file_path)
            elif ext == ".docx":
                preview = extract_preview_from_docx(file_path)
            elif ext == ".pdf":
                preview = extract_preview_from_pdf(file_path)
            elif ext == ".csv":
                preview = extract_preview_from_csv(file_path)
            else:
                preview = ""
            context = guess_file_context(filename, preview)
            context_dict[filename] = context
        except Exception as e:
            context_dict[filename] = f"Lỗi khi lấy ngữ cảnh: {str(e)}"
    if save_json:
        json_path = os.path.join(upload_dir, "file_context.json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(context_dict, f, ensure_ascii=False, indent=2)
    return context_dict

def guess_table_context(table_name, columns):
    prompt = f"Bảng '{table_name}' có các cột: {', '.join(columns)}. Hãy mô tả ngắn ý nghĩa của bảng này bằng tiếng Việt."
    try:
        model = genai.GenerativeModel("gemini-2.0-flash")
        response = model.generate_content(prompt)
        return response.text.strip()
    except Exception as e:
        return f"Không thể sinh mô tả do lỗi API: {str(e)}"

def get_schema_and_context(db_path, save_json=True):
    # Nếu đã có file context -> chỉ đọc, không gọi Gemini nữa
    json_path = os.path.join(os.path.dirname(__file__), "table_context.json")
    if os.path.exists(json_path):
        with open(json_path, "r", encoding="utf-8") as f:
            context_dict = json.load(f)

        # Chuyển về dạng schema_context output cũ
        schema_context = []
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = [row[0] for row in cursor.fetchall()]
        for table in tables:
            cursor.execute(f"PRAGMA table_info({table});")
            columns = [col[1] for col in cursor.fetchall()]
            schema_context.append({
                "table": table,
                "columns": columns,
                "context": context_dict.get(table, "")
            })
        conn.close()

        return schema_context

    # Chưa có file context -> gọi Gemini 1 lần rồi lưu
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = [row[0] for row in cursor.fetchall()]

    schema_context = []
    context_dict = {}

    for table in tables:
        cursor.execute(f"PRAGMA table_info({table});")
        columns = [col[1] for col in cursor.fetchall()]
        try:
            context = guess_table_context(table, columns)
        except Exception as e:
            context = f"LLM error: {str(e)}"

        schema_context.append({
            "table": table,
            "columns": columns,
            "context": context
        })
        context_dict[table] = context

    conn.close()

    if save_json:
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(context_dict, f, ensure_ascii=False, indent=2)

    return schema_context
