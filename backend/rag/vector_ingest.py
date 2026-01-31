import os
import chromadb
from fastapi import APIRouter
from fastapi.responses import JSONResponse
import json
import docx
import openpyxl
import PyPDF2
import csv
from rag.Semantic_for_DB import get_files_context
import uuid
import traceback
import logging
from sentence_transformers import SentenceTransformer
import numpy as np


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
logging.getLogger("chromadb").setLevel(logging.WARNING)
logging.getLogger("chromadb.utils.embedding_functions").setLevel(logging.WARNING)
logging.getLogger("chromadb.utils").setLevel(logging.WARNING)
logging.getLogger("chromadb").propagate = False

router = APIRouter()
UPLOAD_DIR = os.path.join(os.path.dirname(__file__),"..","uploaded_files")
os.makedirs(UPLOAD_DIR, exist_ok=True)


model = SentenceTransformer('all-MiniLM-L6-v2')


class MiniLMEmbeddingFunction:
    def __call__(self, input):
        if isinstance(input, str):
            input = [input]
        embeddings = model.encode(input)
        return [emb.tolist() for emb in embeddings]
    def name(self):
        return "all-MiniLM-L6-v2"
embedding_func = MiniLMEmbeddingFunction()

chroma_client = chromadb.PersistentClient(path=os.path.join(UPLOAD_DIR, "chromadb"))
collection = chroma_client.get_or_create_collection(
    name="rag_vector",
    embedding_function=embedding_func
)

def update_file_context():
    files = [f for f in os.listdir(UPLOAD_DIR) if os.path.splitext(f)[1].lower() in [".xlsx", ".docx", ".pdf", ".txt", ".csv", ".json"]]
    get_files_context(files, UPLOAD_DIR, save_json=True)

def extract_text_from_docx(docx_filename):
    docx_path = os.path.join(UPLOAD_DIR, docx_filename)
    doc = docx.Document(docx_path)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_xlsx(xlsx_filename):
    xlsx_path = os.path.join(UPLOAD_DIR, xlsx_filename)
    wb = openpyxl.load_workbook(xlsx_path)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
    return text

def extract_text_from_pdf(pdf_filename):
    pdf_path = os.path.join(UPLOAD_DIR, pdf_filename)
    text = ""
    with open(pdf_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += page.extract_text() or ""
    return text

def extract_text_from_txt(txt_filename):
    txt_path = os.path.join(UPLOAD_DIR, txt_filename)
    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def extract_text_from_csv(csv_filename):
    csv_path = os.path.join(UPLOAD_DIR, csv_filename)
    texts = []
    with open(csv_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        headers = next(reader, None)
        for row in reader:
            if headers:
               
                row_text = "; ".join([f"{h.strip()}: {v.strip()}" for h, v in zip(headers, row)])
            else:
                row_text = ", ".join(row)
            texts.append(row_text)
    return "\n".join(texts)

def extract_text_from_json(json_filename):
    def flatten_json(y, prefix=""):
        out = []
        if isinstance(y, dict):
            for k, v in y.items():
                out.extend(flatten_json(v, prefix + str(k) + "."))
        elif isinstance(y, list):
            for idx, item in enumerate(y):
                out.extend(flatten_json(item, prefix + f"{idx}."))
        else:
            out.append(f"{prefix[:-1]}: {y}")
        return out

    json_path = os.path.join(UPLOAD_DIR, json_filename)
    with open(json_path, "r", encoding="utf-8", errors="ignore") as f:
        data = json.load(f)
   
    lines = flatten_json(data)
    return "\n".join(lines)

def split_text(text, chunk_size=500, overlap=50):
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:  
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks

def get_context(filename):
    context_path = os.path.join(UPLOAD_DIR, "file_context.json")
    if os.path.isfile(context_path):
        with open(context_path, "r", encoding="utf-8") as f:
            context_data = json.load(f)
        return context_data.get(filename, "")
    return ""

def ingest_file_to_vector_db(
    filename, 
    extract_text_func, 
    collection, 
    context=None, 
    chunk_size=500, 
    overlap=50, 
    batch_size=5
):
    try:
        collection.delete(where={"filename": filename})
        logger.info(f"Đã xóa vectors cũ của file {filename}")
    except Exception as e:
        logger.info(f"Không có vectors cũ để xóa: {str(e)}")
    
    content = extract_text_func(filename)
    if not content.strip():
        logger.warning(f"File {filename} không có nội dung!")
        return

    logger.info(f"Nội dung file {filename}: {len(content)} ký tự")
    chunks = split_text(content, chunk_size=chunk_size, overlap=overlap)
    logger.info(f"Chia thành {len(chunks)} chunks")
    if not chunks:
        logger.warning(f"Không có chunks hợp lệ từ file {filename}")
        return

   
    if chunks:
        logger.info(f"Chunk đầu tiên ({len(chunks[0])} ký tự): {chunks[0][:100]}...")

    successful_chunks = 0
    for i in range(0, len(chunks), batch_size):
        batch_chunks = chunks[i:i+batch_size]
        batch_ids = []
        batch_documents = []
        batch_metadatas = []
        for idx, chunk in enumerate(batch_chunks):
            actual_idx = i + idx
            unique_id = f"{filename}_chunk_{actual_idx}_{uuid.uuid4().hex[:8]}"
            batch_ids.append(unique_id)
            batch_documents.append(chunk)
            batch_metadatas.append({
                "filename": filename,
                "chunk": actual_idx,
                "context": context or ""
            })
        try:
            collection.add(
                documents=batch_documents,
                metadatas=batch_metadatas,
                ids=batch_ids
            )
            successful_chunks += len(batch_chunks)
            logger.info(f"Đã lưu batch {i//batch_size + 1}: {len(batch_chunks)} chunks")
        except Exception as e:
            logger.error(f"Lỗi khi add batch {i//batch_size + 1}: {str(e)}")
            logger.error(traceback.format_exc())
            for j, chunk in enumerate(batch_chunks):
                try:
                    actual_idx = i + j
                    unique_id = f"{filename}_chunk_{actual_idx}_{uuid.uuid4().hex[:8]}"
                    collection.add(
                        documents=[chunk],
                        metadatas=[{"filename": filename, "chunk": actual_idx, "context": context or ""}],
                        ids=[unique_id]
                    )
                    successful_chunks += 1
                    logger.info(f"Đã lưu chunk {actual_idx} riêng lẻ")
                except Exception as chunk_error:
                    logger.error(f"Lỗi chunk {actual_idx}: {str(chunk_error)}")
                    continue
    logger.info(f"Đã lưu {successful_chunks}/{len(chunks)} chunks cho file {filename}")

def check_vectors_by_filename(filename, verbose=False):
    """
    Kiểm tra số lượng vector đã lưu cho một file
    verbose=True để hiển thị chi tiết metadata
    """
    try:
        results = collection.get(where={"filename": filename})
        logger.info(f"File {filename} có {len(results['ids'])} vectors")
        if verbose and len(results['ids']) > 0:
            logger.info(f"Metadata mẫu: {results['metadatas'][0]}")
        return len(results['ids'])
    except Exception as e:
        logger.error(f"Lỗi khi check vectors cho {filename}: {str(e)}")
        return 0

@router.post("/ingest_to_vector/")
async def ingest_to_vector(data: dict):
    filename = data.get("filename")
    try:
        logger.info(f"Bắt đầu ingest file: {filename}")
        ext = os.path.splitext(filename)[1].lower()
        context = get_context(filename)
        if ext == ".txt":
            ingest_file_to_vector_db(filename, extract_text_from_txt, collection, context=context)
        elif ext == ".docx":
            ingest_file_to_vector_db(filename, extract_text_from_docx, collection, context=context)
        elif ext == ".xlsx":
            ingest_file_to_vector_db(filename, extract_text_from_xlsx, collection, context=context)
        elif ext == ".pdf":
            ingest_file_to_vector_db(filename, extract_text_from_pdf, collection, context=context)
        elif ext == ".csv":
            ingest_file_to_vector_db(filename, extract_text_from_csv, collection, context=context)
        elif ext == ".json":
            ingest_file_to_vector_db(filename, extract_text_from_json, collection, context=context)
        else:
            return JSONResponse({"status": "error", "detail": "Unsupported file type"}, status_code=400)
        
        update_file_context()
        
        vector_count = check_vectors_by_filename(filename)
        logger.info(f"Hoàn thành ingest file {filename} với {vector_count} vectors")
        return JSONResponse({
            "status": "success", 
            "filename": filename,
            "vector_count": vector_count
        })
    except Exception as e:
        logger.error(f"Lỗi ingest_to_vector cho file {filename}: {str(e)}")
        logger.error(traceback.format_exc())
        return JSONResponse({"status": "error", "detail": str(e)}, status_code=500)

@router.post("/delete_vector_by_filename/")
async def delete_vector_by_filename(data: dict):
    filename = data.get("filename")
    try:
        results = collection.get(where={"filename": filename})
        logger.info(f"Metadata các vector sẽ xóa: {[m['filename'] for m in results['metadatas']]}")
       
        before_count = check_vectors_by_filename(filename)
        collection.delete(where={"filename": filename})
        after_count = check_vectors_by_filename(filename)
        logger.info(f"Đã xóa {before_count - after_count} vectors cho file {filename}")
        
        context_path = os.path.join(UPLOAD_DIR, "file_context.json")
        if os.path.isfile(context_path):
            with open(context_path, "r", encoding="utf-8") as f:
                context_data = json.load(f)
            if filename in context_data:
                del context_data[filename]
                with open(context_path, "w", encoding="utf-8") as f:
                    json.dump(context_data, f, ensure_ascii=False, indent=2)
        update_file_context()
       
        after_count = check_vectors_by_filename(filename)
        logger.info(f"Đã xóa {before_count - after_count} vectors cho file {filename}")
        return JSONResponse({
            "status": "success",
            "filename": filename,
            "deleted_vectors": before_count - after_count
        })
    except Exception as e:
        logger.error(f"Lỗi delete_vector_by_filename cho file {filename}: {str(e)}")
        return JSONResponse({"status": "error", "detail": str(e)}, status_code=500)